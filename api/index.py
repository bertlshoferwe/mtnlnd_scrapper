"""
Web dashboard for Bid Scout — Vercel + Supabase edition.

This is a stateless Flask app (a single Vercel serverless function). It does
NOT run the scan itself and does NOT schedule anything in-process — Vercel
functions don't stay running between requests, so there's nothing for an
in-process scheduler (APScheduler) to live inside. Scheduling is handled
entirely by .github/workflows/daily-scan.yml (a daily GitHub Actions cron);
this app only reads/writes Supabase and, for "Run Now", asks GitHub to run
that same workflow immediately via its API.

Routes:
  GET  /                                    the dashboard page
  GET  /vendor/<path>                       bundled front-end assets (pdf.js)
  GET  /api/adapters                        list available portal adapters
  GET  /api/divisions                       list divisions
  POST /api/divisions                       create a division
  PATCH /api/divisions/<division_id>        rename a division
  DEL  /api/divisions/<division_id>          delete a division (cascades in Supabase)
  GET  /api/<division_id>/sites              list sites
  POST /api/<division_id>/sites              add a site
  PATCH /api/<division_id>/sites/<site_id>   edit a site (name, url, adapter, selector/pattern)
  POST /api/<division_id>/sites/<site_id>/active   toggle whether the site is scanned
  DEL  /api/<division_id>/sites/<site_id>     remove a site
  GET  /api/<division_id>/keywords           list keywords
  POST /api/<division_id>/keywords           add a keyword
  POST /api/<division_id>/keywords/upload     bulk-add keywords from an uploaded .pdf/.docx (one per line)
  DEL  /api/<division_id>/keywords            remove one keyword (?keyword=…) or all of them
  GET  /api/<division_id>/status             latest run status
  POST /api/<division_id>/cancel-scan        cancel the running scan
  POST /api/<division_id>/run-now            trigger the GitHub Actions workflow now
  GET  /api/<division_id>/results-info       stats + latest run_date + latest summary, for the Results card
  GET  /api/<division_id>/results-grouped    results collapsed to one entry per project, files nested
  GET  /api/<division_id>/results             paginated/filterable rows (search, status, page, page_size) for the Results table
  GET  /download/<division_id>/results        build and stream an .xlsx on the fly from Supabase rows

Every route that touches Supabase or GitHub's API does one or two quick
network calls and returns — well within Vercel Hobby's ~10s function
duration limit, which is exactly why the actual scanning work (which can
easily take minutes) does NOT happen here.
"""

import os
import io
import re
import requests
from datetime import datetime, timezone, timedelta

# A scan_runs row can be left showing status='running' forever if the GitHub
# Actions job that owns it is killed without warning — it hits the workflow's
# timeout-minutes cap, the runner OOMs, or the process gets SIGKILLed — so
# _scan_division's finish_run() never executes. The workflow caps a job at
# 120 min, so any "running" row older than this is definitely dead, not slow;
# the status endpoint reports those as a failed run instead of a live one so
# the dashboard stops showing a phantom "Checking …" indefinitely.
STALE_RUN_AFTER = timedelta(minutes=150)

from flask import Flask, jsonify, request, render_template, send_file, send_from_directory, abort, Response
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
import pypdf
from docx import Document as DocxDocument

import supabase_store
import adapters

_REPO_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
TEMPLATE_DIR = os.path.join(_REPO_ROOT, "templates")
VENDOR_DIR = os.path.join(_REPO_ROOT, "vendor")
app = Flask(__name__, template_folder=TEMPLATE_DIR)

COLUMN_HEADERS = [
    "Date", "Site", "Document URL", "Filename",
    "Matched Keywords", "Match Count", "Keyword Locations", "Status", "AI Notes",
]

# Purely informational — the actual schedule lives in
# .github/workflows/daily-scan.yml. Shown in the dashboard so it's not a
# mystery where/when scans run. Update both places together if you change it.
DISPLAY_SCHEDULE_UTC = os.environ.get("DISPLAY_SCHEDULE_UTC", "07:00 UTC")


def _schedule_hm(text):
    """Pull an "HH:MM" (24h, UTC) out of DISPLAY_SCHEDULE_UTC so the page can
    re-render it in the visitor's local time. Returns None if there's no
    parseable time, in which case the page just shows the raw string."""
    m = re.search(r"(\d{1,2}):(\d{2})", text or "")
    if not m:
        return None
    h, mn = int(m.group(1)), int(m.group(2))
    if 0 <= h < 24 and 0 <= mn < 60:
        return f"{h:02d}:{mn:02d}"
    return None


SCHEDULE_UTC_HM = _schedule_hm(DISPLAY_SCHEDULE_UTC)


def _require_division(division_id):
    division = supabase_store.get_division(division_id)
    if division is None:
        return None, (jsonify({"error": f"no such division '{division_id}'"}), 404)
    return division, None


# ---------------------------------------------------------------------------
# Page
# ---------------------------------------------------------------------------

@app.route("/")
def index():
    return render_template(
        "index.html",
        display_schedule=DISPLAY_SCHEDULE_UTC,
        schedule_utc_hm=SCHEDULE_UTC_HM or "",
    )


@app.route("/vendor/<path:filename>")
def vendor_asset(filename):
    """Serve third-party front-end assets (pdf.js) bundled with the repo.
    Done through Flask rather than Vercel's static handling so it works
    regardless of how the project routes the root path. .mjs needs an explicit
    JS mimetype or browsers refuse to load it as a module."""
    resp = send_from_directory(VENDOR_DIR, filename, max_age=31536000)
    if filename.endswith(".mjs"):
        resp.headers["Content-Type"] = "text/javascript; charset=utf-8"
    resp.headers["Cache-Control"] = "public, max-age=31536000, immutable"
    return resp


# ---------------------------------------------------------------------------
# Divisions
# ---------------------------------------------------------------------------

@app.route("/api/divisions", methods=["GET"])
def api_get_divisions():
    return jsonify(supabase_store.load_divisions())


@app.route("/api/divisions", methods=["POST"])
def api_create_division():
    data = request.get_json(force=True, silent=True) or {}
    try:
        division = supabase_store.create_division(data.get("name", ""))
    except ValueError as e:
        return jsonify({"error": str(e)}), 400
    return jsonify({"ok": True, "division": division})


@app.route("/api/divisions/<division_id>", methods=["PATCH"])
def api_rename_division(division_id):
    data = request.get_json(force=True, silent=True) or {}
    try:
        division = supabase_store.rename_division(division_id, data.get("name", ""))
    except ValueError as e:
        return jsonify({"error": str(e)}), (404 if "not found" in str(e) else 400)
    return jsonify({"ok": True, "division": division})


@app.route("/api/divisions/<division_id>", methods=["DELETE"])
def api_delete_division(division_id):
    try:
        supabase_store.delete_division(division_id)
    except ValueError as e:
        return jsonify({"error": str(e)}), 404
    return jsonify({"ok": True})


# ---------------------------------------------------------------------------
# Sites
# ---------------------------------------------------------------------------

@app.route("/api/<division_id>/sites", methods=["GET"])
def api_get_sites(division_id):
    _, err = _require_division(division_id)
    if err:
        return err
    return jsonify(supabase_store.load_sites(division_id))


def _parse_site_payload(data):
    """Turn the add/edit form's JSON into (name, url, listing, adapter) or
    raise ValueError with a user-facing message. Same rules for both routes
    so a site can be edited into any strategy it could be created in."""
    name = (data.get("name") or "").strip()
    url = (data.get("url") or "").strip()
    adapter = (data.get("adapter") or "").strip()
    if adapter and adapter not in adapters.ADAPTERS:
        raise ValueError(f"unknown adapter '{adapter}'")
    if not name:
        raise ValueError("name is required")
    if not url and not adapter:
        raise ValueError("url is required (unless a portal adapter is selected)")

    # Adapter set -> the adapter knows where to look, listing/pattern ignored.
    # No adapter and no manual selector/pattern -> the generic browser crawler
    # discovers the plan links itself (the "Let AI find the plan links" default).
    # A selector/pattern switches to the targeted listing crawl instead.
    listing = None
    if not adapter and (data.get("link_selector") or data.get("link_pattern")):
        listing = {}
        if data.get("link_selector"):
            listing["link_selector"] = data["link_selector"].strip()
        if data.get("link_pattern"):
            listing["link_pattern"] = data["link_pattern"].strip()

    return name, url, listing, (adapter or None)


@app.route("/api/<division_id>/sites", methods=["POST"])
def api_add_site(division_id):
    _, err = _require_division(division_id)
    if err:
        return err

    try:
        name, url, listing, adapter = _parse_site_payload(
            request.get_json(force=True, silent=True) or {}
        )
    except ValueError as e:
        return jsonify({"error": str(e)}), 400

    site = supabase_store.add_site(
        division_id, name, url, listing=listing, adapter=adapter
    )
    return jsonify({"ok": True, "site": site})


@app.route("/api/<division_id>/sites/<int:site_id>", methods=["PATCH"])
def api_update_site(division_id, site_id):
    _, err = _require_division(division_id)
    if err:
        return err

    try:
        name, url, listing, adapter = _parse_site_payload(
            request.get_json(force=True, silent=True) or {}
        )
    except ValueError as e:
        return jsonify({"error": str(e)}), 400

    try:
        site = supabase_store.update_site(
            division_id, site_id, name, url, listing=listing, adapter=adapter
        )
    except ValueError as e:
        return jsonify({"error": str(e)}), 404
    return jsonify({"ok": True, "site": site})


@app.route("/api/<division_id>/sites/<int:site_id>/active", methods=["POST"])
def api_set_site_active(division_id, site_id):
    """Toggle whether a site is scanned. Body: {"active": true|false}."""
    _, err = _require_division(division_id)
    if err:
        return err
    data = request.get_json(force=True, silent=True) or {}
    active = bool(data.get("active"))
    try:
        site = supabase_store.set_site_active(division_id, site_id, active)
    except ValueError as e:
        msg = str(e)
        return jsonify({"error": msg}), (404 if msg == "site not found" else 400)
    return jsonify({"ok": True, "site": site})


@app.route("/api/adapters", methods=["GET"])
def api_list_adapters():
    return jsonify(adapters.list_adapters())


@app.route("/api/<division_id>/sites/<int:site_id>", methods=["DELETE"])
def api_delete_site(division_id, site_id):
    _, err = _require_division(division_id)
    if err:
        return err
    try:
        supabase_store.delete_site(division_id, site_id)
    except ValueError as e:
        return jsonify({"error": str(e)}), 404
    return jsonify({"ok": True})


# ---------------------------------------------------------------------------
# Keywords
# ---------------------------------------------------------------------------

@app.route("/api/<division_id>/keywords", methods=["GET"])
def api_get_keywords(division_id):
    _, err = _require_division(division_id)
    if err:
        return err
    return jsonify(supabase_store.load_keywords(division_id))


@app.route("/api/<division_id>/keywords", methods=["POST"])
def api_add_keyword(division_id):
    _, err = _require_division(division_id)
    if err:
        return err
    data = request.get_json(force=True, silent=True) or {}
    kw = (data.get("keyword") or "").strip()
    if not kw:
        return jsonify({"error": "keyword is required"}), 400
    try:
        keywords = supabase_store.add_keyword(division_id, kw)
    except ValueError as e:
        return jsonify({"error": str(e)}), 400
    return jsonify({"ok": True, "keywords": keywords})


MAX_KEYWORD_UPLOAD_BYTES = 10 * 1024 * 1024  # 10 MB
MAX_KEYWORD_LENGTH = 120  # skip lines longer than this — likely a sentence/paragraph, not a keyword


def _extract_text_for_keywords(file_bytes, filename):
    """Return the plain text of an uploaded .pdf or .docx, or None for
    anything else. Uses pypdf (not pdfplumber) since this only needs plain
    text, not per-page layout — keeps the Vercel function bundle smaller."""
    lower = filename.lower()
    if lower.endswith(".pdf"):
        reader = pypdf.PdfReader(io.BytesIO(file_bytes))
        return "\n".join((page.extract_text() or "") for page in reader.pages)
    elif lower.endswith(".docx"):
        doc = DocxDocument(io.BytesIO(file_bytes))
        return "\n".join(p.text for p in doc.paragraphs)
    return None


def _lines_to_keywords(text):
    """Turn extracted text into candidate keywords: one per line, trimmed of
    bullets/numbering, blank lines skipped, and anything too long to
    plausibly be a keyword/phrase (rather than a full sentence) dropped."""
    candidates = []
    for line in text.splitlines():
        line = supabase_store.demojibake(line)
        line = line.strip(" \t\u2022\u2023\u25e6\u2043\u2219-–—.").strip()
        line = re.sub(r"^\d+[\.\)]\s*", "", line)  # strip leading "1. " / "2) " numbering
        line = line.strip("*_`").strip()           # strip Markdown emphasis / code ticks
        line = line.rstrip(":").strip()            # strip a trailing "Label:" colon
        if not line or len(line) > MAX_KEYWORD_LENGTH:
            continue
        if supabase_store.looks_like_section_heading(line) or _looks_like_prose(line):
            continue
        candidates.append(line)
    return candidates


def _looks_like_prose(line):
    """Best-effort filter for the explanatory sentences that sit between the
    lists in a structured term document ('These phrases appear in plan
    notes...', 'The most widely specified chamber system...'). Errs toward
    keeping: anything with a brand mark or a parenthetical is left alone."""
    low = line.lower()
    if low.startswith((
        "these ", "the most ", "phrases ", "short strings", "short phrases",
        "compiled for ", "panels are ", "vegetated mse", "sources:",
    )):
        return True
    if "®" in line or "™" in line or "(" in line:  # (R), TM
        return False
    words = line.split()
    return len(words) >= 8 and (
        line.endswith(".") or ". " in line or ";" in line or line.count(",") >= 2
    )


@app.route("/api/<division_id>/keywords/upload", methods=["POST"])
def api_upload_keywords(division_id):
    """Add keywords in bulk from an uploaded .pdf or .docx — one keyword per
    line in the document. Meant for a document that's just a list of terms
    (e.g. a compliance checklist), not for extracting keywords out of prose."""
    _, err = _require_division(division_id)
    if err:
        return err

    if "file" not in request.files or not request.files["file"].filename:
        return jsonify({"error": "choose a .pdf or .docx file first"}), 400

    upload = request.files["file"]
    filename = upload.filename
    file_bytes = upload.read()
    if len(file_bytes) > MAX_KEYWORD_UPLOAD_BYTES:
        return jsonify({"error": "file is too large (10 MB max)"}), 400

    if not filename.lower().endswith((".pdf", ".docx")):
        return jsonify({"error": "only .pdf and .docx files are supported"}), 400

    try:
        text = _extract_text_for_keywords(file_bytes, filename)
    except Exception as e:
        return jsonify({"error": f"couldn't read that file: {e}"}), 400

    if not text or not text.strip():
        return jsonify({"error": "no readable text found in that file"}), 400

    candidates = _lines_to_keywords(text)
    if not candidates:
        return jsonify({"error": "no usable keyword lines found — each keyword should be on its own line"}), 400

    keywords, added, skipped = supabase_store.add_keywords_bulk(division_id, candidates)
    return jsonify({
        "ok": True,
        "keywords": keywords,
        "added": added,
        "skipped_present": skipped["already_present"],
        "skipped_repeat": skipped["repeated_in_file"],
    })


@app.route("/api/<division_id>/keywords", methods=["DELETE"])
def api_delete_keywords(division_id):
    """?keyword=<term> removes that one keyword; with no query param, clears
    every keyword for the division. The term goes in the query string, not
    the path, so slashes / punctuation / ® ™ in keywords survive the round
    trip (a <path:> segment mangles those on some hosts)."""
    _, err = _require_division(division_id)
    if err:
        return err
    keyword = request.args.get("keyword")
    if keyword is not None:
        keywords = supabase_store.delete_keyword(division_id, keyword)
    else:
        keywords = supabase_store.clear_keywords(division_id)
    return jsonify({"ok": True, "keywords": keywords})


# ---------------------------------------------------------------------------
# Status + Run Now (via GitHub Actions workflow_dispatch)
# ---------------------------------------------------------------------------

def _run_is_stale(run):
    """True if a still-'running' scan_runs row is old enough that its GitHub
    Actions job cannot still be alive (see STALE_RUN_AFTER)."""
    started = run.get("started_at")
    if not started:
        return False
    try:
        ts = datetime.fromisoformat(started.replace("Z", "+00:00"))
    except ValueError:
        return False
    if ts.tzinfo is None:
        ts = ts.replace(tzinfo=timezone.utc)
    return datetime.now(timezone.utc) - ts > STALE_RUN_AFTER


@app.route("/api/<division_id>/status", methods=["GET"])
def api_status(division_id):
    _, err = _require_division(division_id)
    if err:
        return err
    run = supabase_store.get_latest_run(division_id)
    if run is None:
        return jsonify({"running": False, "last_started": None, "last_finished": None, "last_result": None})

    running = run["status"] == "running"
    if running and _run_is_stale(run):
        # The owning job is long gone — surface it as a failed run and mark the
        # row so it doesn't keep tripping this check (and so "Run now" isn't
        # blocked by a scan that will never finish).
        running = False
        supabase_store.finish_run(run["id"], "error: timed out")
        run = {**run, "status": "error: timed out",
               "finished_at": datetime.now(timezone.utc).isoformat()}

    return jsonify({
        "running": running,
        "last_started": run["started_at"],
        "last_finished": run["finished_at"],
        "last_result": None if running else run["status"],
        "progress_done": run.get("progress_done") or 0,
        "progress_total": run.get("progress_total") or 0,
        "progress_label": run.get("progress_label") or "",
        "progress_site_i": run.get("progress_site_i") or 0,
        "progress_site_n": run.get("progress_site_n") or 0,
        "progress_overall": run.get("progress_overall") or 0,
    })


def _github_cfg():
    """(headers, api_base, workflow_file, ref) or (None, error_response)."""
    token = os.environ.get("GITHUB_TOKEN")
    owner = os.environ.get("GITHUB_OWNER")
    repo = os.environ.get("GITHUB_REPO")
    if not (token and owner and repo):
        return None, (jsonify({
            "error": "GITHUB_TOKEN/GITHUB_OWNER/GITHUB_REPO not configured on this Vercel "
                     "project — can't talk to the GitHub Actions workflow. See README.md."
        }), 500)
    headers = {
        "Authorization": f"Bearer {token}",
        "Accept": "application/vnd.github+json",
        "X-GitHub-Api-Version": "2022-11-28",
    }
    return (
        headers,
        f"https://api.github.com/repos/{owner}/{repo}",
        os.environ.get("GITHUB_WORKFLOW_FILE", "daily-scan.yml"),
        os.environ.get("GITHUB_REF", "main"),
    ), None


@app.route("/api/<division_id>/run-now", methods=["POST"])
def api_run_now(division_id):
    _, err = _require_division(division_id)
    if err:
        return err

    cfg, cfg_err = _github_cfg()
    if cfg_err:
        return cfg_err
    gh_headers, base, workflow_file, ref = cfg
    owner = base.rsplit("/", 2)[-2]
    repo = base.rsplit("/", 1)[-1]
    url = f"{base}/actions/workflows/{workflow_file}/dispatches"
    try:
        resp = requests.post(
            url, headers=gh_headers,
            json={"ref": ref, "inputs": {"division_id": division_id}},
            timeout=8,
        )
    except requests.RequestException as e:
        return jsonify({"error": f"Failed to reach GitHub: {e}"}), 502

    if resp.status_code < 300:
        return jsonify({"ok": True})

    # Non-2xx: figure out *which* part is wrong so the message is actionable.
    if resp.status_code == 404:
        detail = _diagnose_dispatch_404(base, workflow_file, gh_headers)
        return jsonify({"error": (
            f"GitHub couldn't run the workflow (404). Resolved to "
            f"owner='{owner}', repo='{repo}', workflow='{workflow_file}', ref='{ref}'. {detail}"
        )}), 502

    return jsonify({"error": f"GitHub API error {resp.status_code}: {resp.text}"}), 502


@app.route("/api/<division_id>/cancel-scan", methods=["POST"])
def api_cancel_scan(division_id):
    _, err = _require_division(division_id)
    if err:
        return err

    # Mark this division's running scan as cancelled so the dashboard clears
    # even if the GitHub call below fails or the runner is already gone.
    run = supabase_store.get_latest_run(division_id)
    if run and run.get("status") == "running":
        supabase_store.finish_run(run["id"], "cancelled")

    cfg, cfg_err = _github_cfg()
    if cfg_err:
        return cfg_err
    gh_headers, base, workflow_file, _ = cfg

    # Cancel every in-progress / queued run of the scan workflow. (workflow_
    # dispatch doesn't hand back a run id, and there's usually only one scan
    # in flight, so cancelling all of them is the pragmatic move.)
    cancelled = 0
    try:
        for state in ("in_progress", "queued"):
            r = requests.get(
                f"{base}/actions/workflows/{workflow_file}/runs",
                headers=gh_headers, params={"status": state, "per_page": 20}, timeout=8,
            )
            if not r.ok:
                continue
            for wr in r.json().get("workflow_runs", []):
                c = requests.post(f"{base}/actions/runs/{wr['id']}/cancel",
                                  headers=gh_headers, timeout=8)
                if c.status_code in (202, 409):  # 409 = already completing
                    cancelled += 1
    except requests.RequestException as e:
        return jsonify({"ok": True, "cancelled": cancelled,
                        "warning": f"marked cancelled, but reaching GitHub failed: {e}"})

    return jsonify({"ok": True, "cancelled": cancelled})


def _diagnose_dispatch_404(base, workflow_file, gh_headers):
    """A workflow_dispatch 404 can mean: repo not found / token can't see it,
    the workflow file isn't on the repo, or its name is misspelled. Probe to
    say which."""
    try:
        r = requests.get(base, headers=gh_headers, timeout=6)
        if r.status_code == 404:
            return ("The repo isn't visible to this token — check GITHUB_OWNER/GITHUB_REPO "
                    "and that GITHUB_TOKEN has access (a fine-grained token needs this repo "
                    "selected, with Actions: read & write).")
        if r.status_code == 401:
            return "GITHUB_TOKEN is invalid or expired."
        wr = requests.get(f"{base}/actions/workflows", headers=gh_headers, timeout=6)
        if wr.ok:
            names = sorted(w["path"].split("/")[-1] for w in wr.json().get("workflows", []))
            if workflow_file not in names:
                have = ", ".join(names) or "none"
                return (f"No workflow file named '{workflow_file}' on the default branch. "
                        f"Workflows present: {have}. Set GITHUB_WORKFLOW_FILE to one of those.")
            return ("The workflow exists but the dispatch still 404'd — check GITHUB_REF names a "
                    "real branch and GITHUB_TOKEN has Actions: write.")
    except requests.RequestException:
        pass
    return "Check GITHUB_OWNER, GITHUB_REPO, GITHUB_WORKFLOW_FILE, GITHUB_REF and the token's scopes."


# ---------------------------------------------------------------------------
# Results
# ---------------------------------------------------------------------------

@app.route("/api/<division_id>/results-info", methods=["GET"])
def api_results_info(division_id):
    _, err = _require_division(division_id)
    if err:
        return err
    rows = supabase_store.get_scan_results(division_id, limit=1)
    if not rows:
        return jsonify({"exists": False})
    stats = supabase_store.get_stats(division_id)
    return jsonify({
        "exists": True,
        "latest_run_date": rows[0]["run_date"],
        "stats": stats,
        "latest_summary": supabase_store.get_latest_summary_text(division_id),
    })


@app.route("/api/<division_id>/results", methods=["GET"])
def api_get_results(division_id):
    """Paginated, filterable rows for the dashboard's Results table.
    Query params: search, status, page (1-based), page_size (max 100)."""
    _, err = _require_division(division_id)
    if err:
        return err

    search = (request.args.get("search") or "").strip()
    status = (request.args.get("status") or "").strip()
    try:
        page = max(1, int(request.args.get("page", 1)))
        page_size = max(1, min(100, int(request.args.get("page_size", 20))))
    except ValueError:
        return jsonify({"error": "page and page_size must be integers"}), 400

    rows, total = supabase_store.get_scan_results_page(
        division_id, search=search or None, status=status or None,
        page=page, page_size=page_size,
    )
    return jsonify({"rows": rows, "total": total, "page": page, "page_size": page_size})


@app.route("/api/<division_id>/results-grouped", methods=["GET"])
def api_get_results_grouped(division_id):
    """Results collapsed to one entry per project, files nested. Query params:
    search, status, page (1-based), page_size."""
    _, err = _require_division(division_id)
    if err:
        return err
    search = (request.args.get("search") or "").strip()
    status = (request.args.get("status") or "").strip()
    site = (request.args.get("site") or "").strip()
    keyword = (request.args.get("keyword") or "").strip()
    try:
        page = max(1, int(request.args.get("page", 1)))
        page_size = max(1, min(50, int(request.args.get("page_size", 15))))
    except ValueError:
        return jsonify({"error": "page and page_size must be integers"}), 400

    projects, total, sites = supabase_store.get_results_grouped(
        division_id, search=search or None, status=status or None, site=site or None,
        keyword=keyword or None, page=page, page_size=page_size,
    )
    return jsonify({"projects": projects, "total": total, "sites": sites,
                    "page": page, "page_size": page_size})


_BAD_SHEET_CHARS = re.compile(r"[\[\]:*?/\\]")


def _site_group(site):
    """The site a result belongs to, for grouping into per-site sheets.
    scan_results.site is either "<site name> — <project>" or just the site
    name; either way the part before " — " (or the whole value) is the site."""
    parts = (site or "").split(" — ", 1)
    name = parts[0] if len(parts) == 2 else (site or "")
    return name.strip() or "Other"


def _unique_sheet_name(name, used):
    """A workbook-legal, <=31 char, unique sheet title."""
    base = _BAD_SHEET_CHARS.sub(" ", name).strip()[:31] or "Other"
    candidate, n = base, 2
    while candidate.lower() in used:
        tag = f" ({n})"
        candidate = base[:31 - len(tag)] + tag
        n += 1
    used.add(candidate.lower())
    return candidate


def _write_matches_sheet(ws, rows):
    ws.append(COLUMN_HEADERS)
    for cell in ws[1]:
        cell.font = Font(name="Arial", bold=True)
    widths = [22, 24, 45, 30, 30, 14, 40, 22, 60]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w
    ws.freeze_panes = "A2"
    for r in rows:
        ws.append([
            r["run_date"], r["site"], r["document_url"], r["filename"],
            r["matched_keywords"], r["match_count"], r["keyword_locations"],
            r["status"], r["ai_notes"],
        ])
        for cell in ws[ws.max_row]:
            cell.font = Font(name="Arial")
            cell.alignment = Alignment(wrap_text=True, vertical="top")


def _build_results_workbook(division_id, division_name):
    # The download is a worklist — only documents that hit a keyword. "No
    # match" / "Download failed" / "No documents found" rows stay visible in
    # the dashboard but aren't exported. One sheet per site.
    rows = supabase_store.get_scan_results(division_id, matches_only=True)
    summaries = {s["run_date"]: s["summary"] for s in supabase_store.get_summaries(division_id)}

    groups = {}
    for r in rows:
        groups.setdefault(_site_group(r.get("site")), []).append(r)

    wb = Workbook()
    wb.remove(wb.active)  # drop the default empty sheet

    if not groups:
        _write_matches_sheet(wb.create_sheet("Keyword Matches"), [])
    else:
        used = set()
        for name in sorted(groups, key=str.lower):
            _write_matches_sheet(wb.create_sheet(_unique_sheet_name(name, used)), groups[name])

    if summaries:
        sws = wb.create_sheet("Daily Summary")
        sws.append(["Date", "Summary"])
        for cell in sws[1]:
            cell.font = Font(name="Arial", bold=True)
        sws.column_dimensions["A"].width = 22
        sws.column_dimensions["B"].width = 100
        for run_date in sorted(summaries.keys(), reverse=True):
            sws.append([run_date, summaries[run_date]])
            for cell in sws[sws.max_row]:
                cell.font = Font(name="Arial")
                cell.alignment = Alignment(wrap_text=True, vertical="top")

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


PROXY_MAX_BYTES = 30 * 1024 * 1024


@app.route("/api/<division_id>/pdf", methods=["GET"])
def api_proxy_pdf(division_id):
    """Stream a scanned document through this origin so the dashboard can show
    it in an <iframe> (cross-origin PDFs and X-Frame-Options make direct
    framing unreliable). `?download=1` forces a save instead of inline view.

    SSRF guard: only proxies a URL this division has actually logged in
    scan_results — never an arbitrary URL.
    """
    _, err = _require_division(division_id)
    if err:
        return err

    target = request.args.get("url", "")
    if not target.startswith(("http://", "https://")):
        abort(400, description="bad url")
    if target not in supabase_store.scanned_document_urls_all(division_id):
        abort(403, description="not a document from this division's scans")

    try:
        upstream = requests.get(target, timeout=20, stream=True,
                                headers={"User-Agent": "Mozilla/5.0 (compatible; BidScoutBot/1.0)"})
        upstream.raise_for_status()
    except requests.RequestException as e:
        abort(502, description=f"couldn't fetch the file: {e}")

    filename = _filename_from(target)
    disp = "attachment" if request.args.get("download") else "inline"

    def stream():
        total = 0
        for chunk in upstream.iter_content(65536):
            total += len(chunk)
            if total > PROXY_MAX_BYTES:
                break
            yield chunk

    return Response(stream(), mimetype="application/pdf", headers={
        "Content-Disposition": f'{disp}; filename="{filename}"',
        "Cache-Control": "private, max-age=300",
    })


def _filename_from(url):
    from urllib.parse import unquote, urlparse, parse_qs
    q = parse_qs(urlparse(url).query)
    if "file_name" in q:
        return unquote(q["file_name"][0])
    tail = unquote(urlparse(url).path.rstrip("/").split("/")[-1])
    return tail or "document.pdf"


@app.route("/download/<division_id>/results", methods=["GET"])
def download_results(division_id):
    division, err = _require_division(division_id)
    if err:
        return err

    rows = supabase_store.get_scan_results(division_id, limit=1, matches_only=True)
    if not rows:
        abort(404, description="No keyword matches yet — nothing to download.")

    buf = _build_results_workbook(division_id, division["name"])
    download_name = f"{division['name']}-scan-results.xlsx"
    return send_file(
        buf, as_attachment=True, download_name=download_name,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


# Local testing only — Vercel itself calls `app` directly, never __main__.
if __name__ == "__main__":
    app.run(debug=True, port=5000)
